import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import database

def generate_excel_report(sinf=None, sinf_group=None):
    entries = database.get_entries(sinf=sinf, sinf_group=sinf_group)

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
    
    sub_header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    sub_header_fill = PatternFill(start_color="047857", end_color="047857", fill_type="solid")

    alt_fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")
    regular_font = Font(name="Calibri", size=10)
    
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    # 1. Sheet: Barcha natijalar
    ws_all = wb.create_sheet(title="Barcha natijalar")
    headers_all = [
        "№", "F.I.Sh.", "Sinf", "Guruh", "Test turi", "Til",
        "Ball", "Daraja / Xulosa", "Yetakchi yo'nalishlar", "Fanlar / Qiziqishlar",
        "Kelajak kasbi", "To'garaklar", "Yangi to'garak taklifi", "Takliflar", "Sana va vaqt"
    ]
    ws_all.append(headers_all)
    for col_idx in range(1, len(headers_all) + 1):
        cell = ws_all.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, e in enumerate(entries, start=1):
        row_vals = [
            i,
            e.get('fish', ''),
            e.get('sinf', ''),
            e.get('sinf_group', ''),
            e.get('test_type', ''),
            e.get('til', ''),
            e.get('score', 0),
            e.get('level', ''),
            e.get('dominant_scales', '') or e.get('iqtidor', ''),
            e.get('fanlar', ''),
            e.get('kasb', ''),
            e.get('togaraklar', ''),
            e.get('yangi_togaraklar', ''),
            e.get('takliflar', ''),
            e.get('timestamp', '')
        ]
        ws_all.append(row_vals)
        row_num = ws_all.max_row
        for col_idx in range(1, len(row_vals) + 1):
            cell = ws_all.cell(row=row_num, column=col_idx)
            cell.font = regular_font
            cell.border = thin_border
            if i % 2 == 0:
                cell.fill = alt_fill

    # 2. Sheet: 7-sinf Tashkilotchilik
    entries_7 = [e for e in entries if e.get('test_type') == 'survey_7']
    if entries_7 or sinf_group == '7':
        ws_7 = wb.create_sheet(title="7-sinf Tashkilotchilik")
        headers_7 = ["№", "F.I.Sh.", "Sinf", "Til", "Ball (max 20)", "Daraja", "Xulosa va tavsiya", "Sana"]
        ws_7.append(headers_7)
        for c in range(1, len(headers_7) + 1):
            cell = ws_7.cell(row=1, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for i, e in enumerate(entries_7, start=1):
            row_vals = [
                i, e.get('fish', ''), e.get('sinf', ''), e.get('til', ''),
                e.get('score', 0), e.get('level', ''), e.get('takliflar', ''), e.get('timestamp', '')
            ]
            ws_7.append(row_vals)
            row_num = ws_7.max_row
            for col_idx in range(1, len(row_vals) + 1):
                cell = ws_7.cell(row=row_num, column=col_idx)
                cell.font = regular_font
                cell.border = thin_border
                if i % 2 == 0: cell.fill = alt_fill

    # 3. Sheet: 8-sinf Qobiliyatlar (9 shkala)
    entries_8 = [e for e in entries if e.get('test_type') == 'survey_8']
    if entries_8 or sinf_group == '8':
        ws_8 = wb.create_sheet(title="8-sinf Qobiliyatlar")
        headers_8 = [
            "№", "F.I.Sh.", "Sinf", "Til", "Sport", "Tashkilotchi", "Matematik",
            "Texnika", "Artistik", "Muloqot", "Musiqa", "Rasm", "Adabiyot",
            "Yetakchi qobiliyatlar", "Sana"
        ]
        ws_8.append(headers_8)
        for c in range(1, len(headers_8) + 1):
            cell = ws_8.cell(row=1, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for i, e in enumerate(entries_8, start=1):
            sc = e.get('scale_scores', {})
            def get_sc(key):
                if isinstance(sc, dict) and key in sc:
                    v = sc[key]
                    return v.get('score', 0) if isinstance(v, dict) else v
                return 0
            row_vals = [
                i, e.get('fish', ''), e.get('sinf', ''), e.get('til', ''),
                get_sc('sport'), get_sc('tashkilotchi'), get_sc('matematik'),
                get_sc('texnika'), get_sc('artistic'), get_sc('kommunikativ'),
                get_sc('musiqiy'), get_sc('tasviriy'), get_sc('filologik'),
                e.get('dominant_scales', ''), e.get('timestamp', '')
            ]
            ws_8.append(row_vals)
            row_num = ws_8.max_row
            for col_idx in range(1, len(row_vals) + 1):
                cell = ws_8.cell(row=row_num, column=col_idx)
                cell.font = regular_font
                cell.border = thin_border
                if i % 2 == 0: cell.fill = alt_fill

    # 4. Sheet: 9-sinf Kasbiy xarita (10 shkala)
    entries_9 = [e for e in entries if e.get('test_type') == 'survey_9']
    if entries_9 or sinf_group == '9':
        ws_9 = wb.create_sheet(title="9-sinf Kasbiy xarita")
        headers_9 = [
            "№", "F.I.Sh.", "Sinf", "Til", "Fizika-Matem", "Kimyo-Bio", "Radio-Elektron",
            "Mexanika", "Geografiya", "Adabiyot-San'at", "Tarix-Siyosat", "Pedagogika-Tibbiyot",
            "Biznes-Xo'jalik", "Sport-Harbiy", "Yetakchi kasbiy soha", "Sana"
        ]
        ws_9.append(headers_9)
        for c in range(1, len(headers_9) + 1):
            cell = ws_9.cell(row=1, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for i, e in enumerate(entries_9, start=1):
            sc = e.get('scale_scores', {})
            def get_sc9(key):
                if isinstance(sc, dict) and key in sc:
                    v = sc[key]
                    return v.get('score', 0) if isinstance(v, dict) else v
                return 0
            row_vals = [
                i, e.get('fish', ''), e.get('sinf', ''), e.get('til', ''),
                get_sc9('fizika_matematika'), get_sc9('kimyo_biologiya'), get_sc9('radiotexnika'),
                get_sc9('mexanika'), get_sc9('geografiya'), get_sc9('adabiyot_sanat'),
                get_sc9('tarix_siyosat'), get_sc9('pedagogika_tibbiyot'), get_sc9('tadbirkorlik'),
                get_sc9('sport_harbiy'), e.get('dominant_scales', ''), e.get('timestamp', '')
            ]
            ws_9.append(row_vals)
            row_num = ws_9.max_row
            for col_idx in range(1, len(row_vals) + 1):
                cell = ws_9.cell(row=row_num, column=col_idx)
                cell.font = regular_font
                cell.border = thin_border
                if i % 2 == 0: cell.fill = alt_fill

    # Adjust column widths for all sheets
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 45)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
